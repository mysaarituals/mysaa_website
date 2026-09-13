"""
Converts data/mysaa_products.xlsx back into the JSON files the website
reads (products.json, fragrances.json, categories.json, settings.json).

Run this after editing the Excel workbook:

    python3 scripts/xlsx_to_json.py

Requires: pip install openpyxl
"""
import json
import os
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
XLSX = os.path.join(DATA, "mysaa_products.xlsx")


def rows_as_dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() for h in rows[0]]
    out = []
    for row in rows[1:]:
        if row[0] in (None, ""):
            continue
        out.append(dict(zip(headers, row)))
    return out


def as_bool(v):
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    return str(v).strip().lower() in ("true", "1", "yes")


def main():
    wb = load_workbook(XLSX, data_only=True)

    products = []
    for r in rows_as_dicts(wb["Products"]):
        products.append({
            "slug": r.get("slug", ""),
            "name": r.get("name", ""),
            "categorySlug": r.get("categorySlug", ""),
            "fragranceSlug": r.get("fragranceSlug", ""),
            "price": r.get("price", 0) or 0,
            "volume": r.get("volume", "") or "",
            "weight": r.get("weight", "") or "",
            "materials": r.get("materials", "") or "",
            "packaging": r.get("packaging", "") or "",
            "shortDescription": r.get("shortDescription", "") or "",
            "about": r.get("about", "") or "",
            "bestseller": as_bool(r.get("bestseller")),
            "isNew": as_bool(r.get("isNew")),
            "customizable": as_bool(r.get("customizable")),
            "active": as_bool(r.get("active")),
            "images": [],
        })

    fragrances = [{
        "slug": r.get("slug", ""), "name": r.get("name", ""), "notes": r.get("notes", "") or "",
        "description": r.get("description", "") or "", "mood": r.get("mood", "") or "",
        "active": as_bool(r.get("active")),
    } for r in rows_as_dicts(wb["Fragrances"])]

    categories = [{
        "slug": r.get("slug", ""), "name": r.get("name", ""),
        "subtitle": r.get("subtitle", "") or "", "active": as_bool(r.get("active")),
    } for r in rows_as_dicts(wb["Categories"])]

    settings = {}
    for row in wb["Settings"].iter_rows(min_row=2, values_only=True):
        if row[0]:
            settings[str(row[0])] = row[1]

    def dump(name, data):
        path = os.path.join(DATA, f"{name}.json")
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
        print(f"Wrote {path}")

    dump("products", products)
    dump("fragrances", fragrances)
    dump("categories", categories)
    dump("settings", settings)


if __name__ == "__main__":
    main()
