"""
Builds data/mysaa_products.xlsx from the JSON files in /data.
Run this once to generate the Excel "master" workbook that the
shop owner can open and edit in Excel / Google Sheets.

    python3 scripts/build_xlsx_from_json.py
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

HEADER_FILL = PatternFill(start_color="7A1420", end_color="7A1420", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def load(name):
    with open(os.path.join(DATA, f"{name}.json"), encoding="utf-8") as f:
        return json.load(f)


def write_sheet(ws, headers, rows):
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 12), 50)
    ws.freeze_panes = "A2"


def main():
    products = load("products")
    fragrances = load("fragrances")
    feelings = load("feelings")
    occasions = load("occasions")
    categories = load("categories")
    settings = load("settings")

    wb = Workbook()

    ws = wb.active
    ws.title = "Products"
    headers = [
        "slug", "name", "type", "categorySlug", "fragranceSlugs", "feelingSlugs",
        "occasionSlugs", "price", "priceOnRequest", "size", "burnTime", "materials",
        "packaging", "shortDescription", "ritual_fragrance", "ritual_feeling",
        "ritual_inspiration", "ritual_perfectFor", "featured", "bestseller",
        "isNew", "customizable", "active",
    ]
    rows = []
    for p in products:
        r = p.get("ritual", {})
        rows.append([
            p.get("slug", ""), p.get("name", ""), p.get("type", ""), p.get("categorySlug", ""),
            ", ".join(p.get("fragranceSlugs", [])), ", ".join(p.get("feelingSlugs", [])),
            ", ".join(p.get("occasionSlugs", [])), p.get("price", 0), p.get("priceOnRequest", False),
            p.get("size", ""), p.get("burnTime", ""), p.get("materials", ""), p.get("packaging", ""),
            p.get("shortDescription", ""), r.get("fragrance", ""), r.get("feeling", ""),
            r.get("inspiration", ""), r.get("perfectFor", ""), p.get("featured", False),
            p.get("bestseller", False), p.get("isNew", False), p.get("customizable", False),
            p.get("active", True),
        ])
    write_sheet(ws, headers, rows)

    ws2 = wb.create_sheet("Fragrances")
    write_sheet(ws2, ["slug", "name", "notes", "description", "mood", "active"],
                [[f["slug"], f["name"], f["notes"], f["description"], f["mood"], f["active"]] for f in fragrances])

    ws3 = wb.create_sheet("Feelings")
    write_sheet(ws3, ["slug", "name", "description", "active"],
                [[f["slug"], f["name"], f["description"], f["active"]] for f in feelings])

    ws4 = wb.create_sheet("Occasions")
    write_sheet(ws4, ["slug", "name", "active"],
                [[o["slug"], o["name"], o["active"]] for o in occasions])

    ws5 = wb.create_sheet("Categories")
    write_sheet(ws5, ["slug", "name", "active"],
                [[c["slug"], c["name"], c["active"]] for c in categories])

    ws6 = wb.create_sheet("Settings")
    write_sheet(ws6, ["key", "value"], [[k, v] for k, v in settings.items()])

    ws0 = wb.create_sheet("Read Me", 0)
    instructions = [
        ["MYSAA RITUALS — Product Data Workbook"],
        [""],
        ["This workbook is the master source for everything shown on the website."],
        ["Edit any sheet below, save the file, then run the conversion script to"],
        ["turn your changes back into the JSON files the website reads."],
        [""],
        ["1. Edit 'Products', 'Fragrances', 'Feelings', 'Occasions', 'Categories'"],
        ["   or 'Settings' below. Keep the 'slug' column unique on each sheet —"],
        ["   it is how products link to fragrances, feelings and occasions."],
        ["2. In 'Products', list multiple fragranceSlugs / feelingSlugs / occasionSlugs"],
        ["   separated by a comma and a space, e.g.  dhoop-chandan, madhuban"],
        ["3. Save this file as mysaa_products.xlsx in the data/ folder."],
        ["4. From a terminal in the project folder, run:"],
        ["      python3 scripts/xlsx_to_json.py"],
        ["   This regenerates the .json files the site uses."],
        ["5. Commit and push the changed files in /data to GitHub. GitHub Pages"],
        ["   will update the live site automatically within a minute or two."],
        [""],
        ["No coding needed for steps 1-3 — only step 4 needs a terminal."],
    ]
    for row in instructions:
        ws0.append(row)
    ws0["A1"].font = Font(bold=True, size=14, color="7A1420")
    ws0.column_dimensions["A"].width = 90

    out_path = os.path.join(DATA, "mysaa_products.xlsx")
    wb.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
